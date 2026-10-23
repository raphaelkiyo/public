"""MITERAS：手動ログイン・ローカル出力。Python 3.10以上。"""
import argparse
import builtins
import json
import math
import os
import sys
import time
import unicodedata
from collections import defaultdict
from copy import copy
from datetime import date, datetime, time as dt_time, timedelta, timezone
from pathlib import Path

JST = timezone(timedelta(hours=9))
LOGIN_URL = 'https://kintai.miteras.jp/PXT_PTCS/login'
BASE = 'https://kintai.miteras.jp/PXT_PTCS/'


def profile_directory(channel):
    """通常のChromeとは別の、ユーザー専用プロファイルを使用します。"""
    if sys.platform == 'win32':
        root = Path(os.environ.get('LOCALAPPDATA') or Path.home() / 'AppData' / 'Local')
    else:
        root = Path(os.environ.get('XDG_STATE_HOME') or Path.home() / '.local' / 'state')
    return root / 'MiterasExporter' / f'{channel}-profile'


def stamp(day):
    return datetime.combine(day, dt_time(), JST).astimezone(timezone.utc).isoformat(timespec='milliseconds').replace('+00:00', 'Z')


def local_date(ms):
    if isinstance(ms, bool) or not isinstance(ms, (int, float)):
        raise ValueError('応答の日付が未設定、または不正です。')
    return datetime.fromtimestamp(ms / 1000, JST).date()


def minutes(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0 or value != int(value):
        raise ValueError('時間が未設定、負数、または整数の分数ではありません。')
    return int(value)


def clock_value(ms, day):
    """対象日の午前0時からの経過時間。翌日1時は25:00と表示します。"""
    if ms is None:
        return None
    if isinstance(ms, bool) or not isinstance(ms, (int, float)) or not math.isfinite(ms):
        raise ValueError('始業・終業時刻の形式が不正です。')
    elapsed = (datetime.fromtimestamp(ms / 1000, JST) - datetime.combine(day, dt_time(), JST)).total_seconds()
    if elapsed < 0:
        raise ValueError('始業・終業時刻が対象日より前になっています。')
    return elapsed / 86400


def signed_duration(value):
    """Excelで負の時刻が####にならないよう、差分は符号付き文字列にします。"""
    if value is None:
        return None
    sign = '-' if value < 0 else '+' if value > 0 else ''
    return f'{sign}{abs(value) // 60}:{abs(value) % 60:02d}'


def parse_day(day, attendance, projects):
    daily = attendance.get('dailyAttendance')
    if not isinstance(daily, dict):
        raise ValueError('dailyAttendanceが未設定のため、対象日のデータを確認できません。')
    if local_date(daily.get('workDate')) != day:
        raise ValueError('応答の日付が指定日と一致しません。')
    details = projects.get('projectTaskDetail')
    if not isinstance(details, list):
        raise ValueError('projectTaskDetailが未設定またはnullです。空の一覧とは区別します。')
    refs = projects.get('projectTaskReference') or []
    catalog = {}
    for ref in refs:
        project = ref.get('project') or {}
        catalog[project.get('projectCode')] = project.get('projectDescription') or project.get('projectName') or ''
    rows, notes, seen = [], [], set()
    running_total = 0
    for detail in details:
        code = detail.get('projectCode')
        if not isinstance(code, str) or not code:
            raise ValueError('プロジェクトコードがありません。')
        logs = detail.get('loggedHours')
        if not isinstance(logs, list):
            raise ValueError('loggedHoursが未設定またはnullです。')
        subtotal = 0
        for log in logs:
            if local_date(log.get('taskDate')) != day:
                raise ValueError('工数明細の日付が指定日と一致しません。')
            if log.get('projectCode') != code:
                raise ValueError('明細のプロジェクトコードが一致しません。')
            serial = log.get('serialNumber')
            if serial is None:
                raise ValueError('明細にserialNumberがありません。')
            key = (code, serial)
            if key in seen:
                raise ValueError('応答に重複明細があります。二重計上を防ぐため、この日の集計を中止しました。')
            seen.add(key)
            effort = minutes(log.get('effort'))
            subtotal += effort
            rows.append([day, code, catalog.get(code, ''), serial, log.get('taskId'), effort,
                         effort / 1440, log.get('remarks') or '',
                         log.get('approvalCondition'), log.get('confirmationCondition')])
        running_total += subtotal
        reference_total = detail.get('totalProjectTime')
        if reference_total is not None:
            reference_total = minutes(reference_total)
            # 応答にはプロジェクト単独の合計と、先頭からの累計の両形式があります。
            if reference_total not in (subtotal, running_total):
                notes.append(f'{code}: 参考合計{reference_total}分が明細合計・累計と一致しません')
        if not catalog.get(code):
            notes.append(f'{code}: プロジェクト内容が見つかりません')
    worked = daily.get('workTime')
    worked = None if worked is None else minutes(worked)
    total = sum(r[5] for r in rows)
    difference = None if worked is None else total - worked
    work_type = daily.get('workTypeName') or daily.get('scheduleLeaveName') or (daily.get('shift') or {}).get('leaveName') or ''
    is_rest = (work_type in ('所定休日', '法定休日', '休日') and not rows
               and worked in (None, 0) and daily.get('workStartTime') is None
               and daily.get('workEndTime') is None and daily.get('breakTime') in (None, 0))
    if worked is None and not is_rest:
        notes.append('workTimeが未設定のため、日別合計を照合できません')
    elif difference:
        notes.append('工数合計とworkTimeが一致しません。MITERASで確認してください')
    status = '要確認' if notes else ('休日' if is_rest else '正常' if rows else '登録なし')
    start_time = clock_value(daily.get('workStartTime'), day)
    end_time = clock_value(daily.get('workEndTime'), day)
    break_minutes = daily.get('breakTime')
    break_duration = None if break_minutes is None else minutes(break_minutes) / 1440
    return rows, [day, status, worked, total, difference, worked / 1440 if worked is not None else None,
                  total / 1440, '; '.join(notes), work_type, start_time, end_time, break_duration]


class ExcelTextMetrics:
    """Windowsでは実際の日本語フォントで幅を測定します。追加ライブラリ不要。"""
    def __init__(self):
        self.dc = None
        self.fonts = {}
        self.cache = {}
        if sys.platform != 'win32':
            return
        try:
            import ctypes
            from ctypes import wintypes
            self.ctypes = ctypes
            self.size_type = wintypes.SIZE
            self.gdi = ctypes.WinDLL('gdi32')
            self.gdi.CreateCompatibleDC.argtypes = [wintypes.HDC]
            self.gdi.CreateCompatibleDC.restype = wintypes.HDC
            self.gdi.CreateFontW.argtypes = [ctypes.c_int] * 5 + [wintypes.DWORD] * 8 + [wintypes.LPCWSTR]
            self.gdi.CreateFontW.restype = wintypes.HFONT
            self.gdi.SelectObject.argtypes = [wintypes.HDC, wintypes.HGDIOBJ]
            self.gdi.SelectObject.restype = wintypes.HGDIOBJ
            self.gdi.GetTextExtentPoint32W.argtypes = [wintypes.HDC, wintypes.LPCWSTR, ctypes.c_int, ctypes.POINTER(wintypes.SIZE)]
            self.gdi.GetTextExtentPoint32W.restype = wintypes.BOOL
            self.gdi.DeleteObject.argtypes = [wintypes.HGDIOBJ]
            self.gdi.DeleteDC.argtypes = [wintypes.HDC]
            self.dc = self.gdi.CreateCompatibleDC(None)
            if not self.dc:
                return
            for bold in (False, True):
                self.fonts[bold] = self.gdi.CreateFontW(-15, 0, 0, 0, 700 if bold else 400,
                                                     0, 0, 0, 128, 0, 0, 4, 0, 'Yu Gothic')
            self.original_font = self.gdi.SelectObject(self.dc, self.fonts[False])
        except Exception:
            self.close()

    def width(self, text, bold=False):
        text = str(text or '').replace('\t', '    ')
        key = (text, bool(bold))
        if key in self.cache:
            return self.cache[key]
        width = 0
        for line in text.split('\n'):
            pixels = None
            if self.dc and self.fonts.get(bool(bold)):
                self.gdi.SelectObject(self.dc, self.fonts[bool(bold)])
                size = self.size_type()
                if self.gdi.GetTextExtentPoint32W(self.dc, line, len(line.encode('utf-16-le')) // 2, self.ctypes.byref(size)):
                    pixels = size.cx
            if pixels is None:
                pixels = sum(15 if unicodedata.east_asian_width(ch) in 'WF' else 7.5 for ch in line)
                if bold:
                    pixels *= 1.06
            width = max(width, pixels)
        self.cache[key] = width
        return width

    def close(self):
        if self.dc:
            if getattr(self, 'original_font', None):
                self.gdi.SelectObject(self.dc, self.original_font)
            for font in self.fonts.values():
                if font:
                    self.gdi.DeleteObject(font)
            self.gdi.DeleteDC(self.dc)
        self.dc = None


def write_excel(path, rows, days, start, end):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    book = Workbook()
    detail = book.active
    detail.title = '工数明細'
    summary = book.create_sheet('プロジェクト別集計')
    check = book.create_sheet('日別照合')
    summary.append(['プロジェクトコード', 'プロジェクト内容', '合計工数（h）'])
    agg = defaultdict(lambda: {'minutes': 0, 'days': set(), 'description': ''})
    for r in rows:
        a = agg[r[1]]
        a['minutes'] += r[5]
        a['days'].add(r[0])
        a['description'] = r[2] or a['description']
    for code, a in sorted(agg.items()):
        summary.append([code, a['description'], a['minutes'] / 60])
    detail.append(['日付', 'プロジェクトコード', 'プロジェクト内容', '工数（h）', '備考',
                   '勤務（h）', '工数合計（h）', '差分（h）', '照合結果', '確認事項'])
    grouped = defaultdict(list)
    for row in rows:
        grouped[row[0]].append(row)
    by_day = {row[0]: row for row in days}
    detail_groups = []
    for group_index, day in enumerate(sorted(set(grouped) | set(by_day))):
        records = grouped.get(day, [])
        daily = by_day.get(day)
        if daily is None:
            total_minutes = sum(row[5] for row in records)
            daily = [day, '要確認', None, total_minutes, None, None, total_minutes / 1440,
                     '日別勤怠データがないため照合できません。']
        first_row = detail.max_row + 1
        for index, row in enumerate(records):
            detail.append([day if index == 0 else None, row[1], row[2], row[5] / 60, row[7]])
        if len(records) > 1:
            # 工数列には日計を重複記載せず、右側の工数合計列に表示します。
            detail.append([None, None, '日計', None, None])
        elif not records:
            work_type = daily[8] if len(daily) > 8 else ''
            label = work_type if daily[1] == '休日' else '取得エラー' if daily[1] == 'エラー' else '工数未登録'
            detail.append([day, None, label, None if daily[1] in ('休日', 'エラー') else 0, None])
        last_row = detail.max_row
        worked_hours = None if daily[2] is None else daily[2] / 60
        project_hours = None if daily[3] is None else daily[3] / 60
        difference_hours = None if daily[4] is None else daily[4] / 60
        for col, value in enumerate([worked_hours, project_hours, difference_hours, daily[1], daily[7] or None], 6):
            detail.cell(last_row, col, value)
        detail_groups.append((first_row, last_row, group_index, len(records) > 1, daily[1]))
    check.append(['日付', '勤務区分', '始業時刻', '終業時刻', '休憩時間（h）',
                  '勤務時間（h）', '工数合計（h）', '差分（工数－勤務・h）', '照合結果', '詳細'])
    for row in days:
        extra = row[8:12] if len(row) >= 12 else [None] * 4
        break_hours = None if extra[3] is None else extra[3] * 24
        worked_hours = None if row[2] is None else row[2] / 60
        project_hours = None if row[3] is None else row[3] / 60
        difference_hours = None if row[4] is None else row[4] / 60
        check.append([row[0], *extra[:3], break_hours, worked_hours, project_hours, difference_hours, row[1], row[7]])
    for ws in book.worksheets:
        ws.freeze_panes = 'A2'
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(name='Yu Gothic', color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='193B59')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Yu Gothic', size=11)
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                if isinstance(cell.value, str):
                    # 備考が次の文字で始まる場合も、数式ではなく文字列として保存します： '='.
                    cell.data_type = 's'
                if isinstance(cell.value, date):
                    cell.number_format = 'yyyy/mm/dd'
    for ws, cols in [(summary, [3]), (detail, [4, 6, 7, 8]), (check, [5, 6, 7, 8])]:
        for row in range(2, ws.max_row + 1):
            for col in cols:
                ws.cell(row, col).number_format = '0.00'
    for row in range(2, check.max_row + 1):
        for col in (3, 4):
            check.cell(row, col).number_format = '[h]:mm'
    grid = Side(style='thin', color='404040')
    separator = Side(style='thin', color='A9B7C4')
    subtotal_edge = Side(style='thin', color='D6DEE6')
    for ws in (summary, check):
        for row in ws:
            for cell in row:
                cell.border = Border(left=grid, right=grid, top=grid, bottom=grid)
    for first, last, group_index, has_subtotal, status in detail_groups:
        fill = PatternFill('solid', fgColor='F3F4F5' if status == '休日' else 'F3F7FB' if group_index % 2 == 0 else 'FFFFFF')
        for row in detail.iter_rows(min_row=first, max_row=last):
            for cell in row:
                cell.fill = fill
                cell.border = Border(top=separator if cell.row == first else Side(),
                                     bottom=separator if cell.row == last else Side(),
                                     right=subtotal_edge if cell.column == 5 else Side())
                if status == '休日':
                    cell.font = Font(name='Yu Gothic', size=11, color='777777')
        detail.cell(first, 1).font = Font(name='Yu Gothic', size=11, bold=True, color='193B59')
        if has_subtotal:
            for cell in detail[last]:
                cell.fill = fill
                cell.font = Font(name='Yu Gothic', size=11, bold=True, color='193B59')
                cell.border = Border(top=subtotal_edge, bottom=separator,
                                     right=subtotal_edge if cell.column == 5 else Side())
        else:
            for col in (6, 7):
                detail.cell(last, col).font = Font(name='Yu Gothic', size=11, bold=True, color='193B59')
        if status in ('要確認', 'エラー'):
            for col in (8, 9):
                detail.cell(last, col).fill = PatternFill('solid', fgColor='FFF0CC')
    for row in check.iter_rows(min_row=2):
        if row[8].value in ('エラー', '要確認'):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='FFF0CC')
    # 全見出しを中央揃え。指定の文章列のみ、データを左揃えにします。
    left_aligned_headers = {'プロジェクト内容', '確認事項', '詳細', '備考'}
    for ws in book.worksheets:
        for row in ws:
            for cell in row:
                header = ws.cell(1, cell.column).value
                cell.alignment = Alignment(
                    horizontal='left' if cell.row > 1 and header in left_aligned_headers else 'center',
                    vertical='center', wrap_text=cell.row > 1 and header in ('確認事項', '詳細'))
    # フォントの実測幅から列幅を調整します。長い確認事項のみ折り返します。
    metrics = ExcelTextMetrics()
    def displayed(cell):
        value = cell.value
        if value is None:
            return ''
        if isinstance(value, date):
            return value.strftime('%Y/%m/%d')
        if isinstance(value, (int, float)):
            if cell.number_format == '[h]:mm':
                total_minutes = round(value * 1440)
                return f'{total_minutes // 60}:{total_minutes % 60:02d}'
            return f'{value:.2f}'
        return str(value)
    try:
        for ws in book.worksheets:
            for column in ws.columns:
                width_px = max(metrics.width(displayed(cell), cell.font.bold) for cell in column)
                width = max(10, (width_px + 18) / 7)
                # 詳細・確認事項は読みやすい幅を上限とし、行高で全内容を表示します。
                if ws in (detail, check) and column[0].column == 10:
                    width = min(55, width)
                ws.column_dimensions[column[0].column_letter].width = min(255, round(width, 2))
                ws.column_dimensions[column[0].column_letter].bestFit = True
            for row in ws:
                line_count = 1
                for cell in row:
                    text = displayed(cell)
                    available = max(1, ws.column_dimensions[cell.column_letter].width * 7 - 12)
                    lines = text.split('\n')
                    if cell.alignment.wrap_text:
                        count = sum(max(1, math.ceil(metrics.width(line, cell.font.bold) / available)) for line in lines)
                    else:
                        count = len(lines)
                    line_count = max(line_count, count)
                ws.row_dimensions[row[0].row].height = 18 * line_count + (10 if row[0].row == 1 else 6)
    finally:
        metrics.close()
    for first, last, _, _, _ in detail_groups:
        if last > first:
            block_fill = copy(detail.cell(first, 1).fill)
            detail.merge_cells(start_row=first, end_row=last, start_column=1, end_column=1)
            for row in range(first, last + 1):
                detail.cell(row, 1).fill = block_fill
            detail.cell(first, 1).alignment = Alignment(horizontal='center', vertical='center')
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.tmp.xlsx')
    book.save(temp)
    temp.replace(path)


def run(start, end, output, report=None, login_wait=None, cancelled=None, progress=None, saved=None):
    from playwright.sync_api import sync_playwright
    print = report or builtins.print
    rows, days = [], []
    state = {'csrf': None}
    def observe(request):
        if request.url.startswith(BASE):
            token = request.all_headers().get('x-csrf-token')
            if token:
                state['csrf'] = token
    with sync_playwright() as p:
        context = None
        for channel in ('chrome', 'msedge'):
            try:
                profile = profile_directory(channel)
                profile.mkdir(parents=True, exist_ok=True)
                context = p.chromium.launch_persistent_context(
                    user_data_dir=str(profile), channel=channel, headless=False,
                    chromium_sandbox=True, timezone_id='Asia/Tokyo', locale='ja-JP')
                print(f'ブラウザー：{channel}')
                break
            except Exception:
                pass
        if context is None:
            raise RuntimeError('ChromeまたはEdgeを起動できません。インストール状況を確認し、このアプリの別ウィンドウを閉じてください。')
        context.on('request', observe)
        page = context.pages[0] if context.pages else context.new_page()
        try:
            page.goto(LOGIN_URL, wait_until='domcontentloaded', timeout=60000)
            print('\n専用プロファイルのセッションを再利用します。ログイン画面が表示された場合はログインしてください。')
            print('対象日の詳細とプロジェクトタブを開いてから、Excel出力に進んでください。')
            if login_wait:
                login_wait(page)
            else:
                input('この画面に戻り、ENTERキーを押すと取得を開始します：')
            page.wait_for_timeout(300)
            if not state['csrf']:
                for tab in context.pages:
                    if tab.url.startswith(BASE):
                        meta = tab.locator('meta[name="_csrf"]')
                        if meta.count():
                            state['csrf'] = meta.first.get_attribute('content')
                            break
            if not state['csrf']:
                raise RuntimeError('CSRFトークンを検出できません。再実行し、ログイン後に日別詳細とプロジェクトタブを開いてください。')
            def query(endpoint, day):
                for attempt in range(3):
                    response = context.request.post(BASE + endpoint,
                        data=json.dumps(stamp(day)), timeout=30000,
                        headers={'Content-Type': 'application/json', 'Accept': 'application/json',
                                 'X-CSRF-TOKEN': state['csrf'], 'X-Requested-With': 'XMLHttpRequest',
                                 'Client-Timezone-Offset': '-540', 'Origin': BASE.split('/PXT_')[0],
                                 'Referer': BASE + 'work-condition'})
                    try:
                        if response.status in (401, 403):
                            raise PermissionError('セッションの有効期限が切れたか、アクセスが拒否されました。再度ログインしてください。')
                        if response.status == 429 or 500 <= response.status <= 599:
                            if attempt < 2:
                                time.sleep(1 + attempt * 2)
                                continue
                        if response.status != 200:
                            raise ValueError(f'取得時にHTTP {response.status}が返されました。')
                        if 'json' not in response.headers.get('content-type', '').lower():
                            raise PermissionError('応答がJSONではありません。ログイン画面に転送された可能性があります。')
                        try:
                            result = response.json()
                        except Exception:
                            raise ValueError('不正なJSONを受信しました。') from None
                        if not isinstance(result, dict):
                            raise ValueError('JSON応答の形式が想定と異なります。')
                        return result
                    finally:
                        response.dispose()
                raise ValueError('データ取得が完了しませんでした。')
            day = start
            stop = False
            while day <= end:
                if cancelled and cancelled.is_set():
                    raise KeyboardInterrupt
                print(f'{day:%Y/%m/%d}を取得中… ', end='', flush=True)
                if stop:
                    days.append([day, 'エラー', None, None, None, None, None, '認証エラーが発生したため、取得していません。'])
                    print('未取得')
                else:
                    try:
                        attendance = query('retrieveDailyAttendance', day)
                        projects = query('retrieveDailyAttendanceProjectTasks', day)
                        day_rows, check = parse_day(day, attendance, projects)
                        rows.extend(day_rows)
                        days.append(check)
                        print(check[1])
                    except PermissionError as exc:
                        days.append([day, 'エラー', None, None, None, None, None, str(exc)])
                        stop = True
                        print('認証エラー')
                    except ValueError as exc:
                        days.append([day, 'エラー', None, None, None, None, None, str(exc)])
                        print('データ・取得エラー')
                    except Exception:
                        # 認証ヘッダーを含む可能性があるため、通信例外の内容は表示しません。
                        days.append([day, 'エラー', None, None, None, None, None, '通信エラー、または想定外のデータ形式です。'])
                        print('通信・形式エラー')
                    time.sleep(0.25)
                if progress:
                    progress((day - start).days + 1, (end - start).days + 1, day, days[-1][1])
                day += timedelta(days=1)
        finally:
            context.close()
            state.clear()
    incomplete = any(d[1] == 'エラー' for d in days)
    if incomplete:
        output = output.with_name(output.stem + '_未完了' + output.suffix)
    if progress:
        progress((end - start).days + 1, (end - start).days + 1, None, '保存中')
    write_excel(output, rows, days, start, end)
    if saved:
        saved(output)
    print(f'\nExcel: {output.resolve()}')
    print('「日別照合」シートを確認してください。' + (' エラーの日があるため、合計は取得済み分のみです。' if incomplete else ''))
    return 2 if incomplete else 0


def main():
    parser = argparse.ArgumentParser(description='MITERAS → Excel（手動ログイン）')
    parser.add_argument('--inicio', help='日付（対象日を含む）：YYYY-MM-DD')
    parser.add_argument('--fim', help='日付（対象日を含む）：YYYY-MM-DD')
    parser.add_argument('--saida', type=Path, help='出力先の.xlsxファイル')
    args = parser.parse_args()
    try:
        start = date.fromisoformat(args.inicio or input('開始日（YYYY-MM-DD）：').strip())
        end = date.fromisoformat(args.fim or input('終了日（YYYY-MM-DD）：').strip())
        if end < start:
            raise ValueError('終了日が開始日より前になっています。')
        if (end - start).days > 365:
            raise ValueError('1回の取得期間は最大366日です。')
        output = args.saida or Path(__file__).resolve().parent / f'miteras_{start}_{end}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx'
        if output.suffix.lower() != '.xlsx':
            raise ValueError('出力ファイルの拡張子は.xlsxにしてください。')
        if output.exists() or output.with_name(output.stem + '_未完了.xlsx').exists():
            raise ValueError('同名のファイルが存在します。以前の帳票を残すため、別の名前を指定してください。')
        return run(start, end, output)
    except (ValueError, RuntimeError, ImportError) as exc:
        print(f'エラー：{exc}')
        return 1
    except KeyboardInterrupt:
        print('\nユーザーによって中止されました。')
        return 130
    except Exception:
        print('ブラウザーの起動または帳票の保存に失敗しました。環境を確認し、保存先のExcelファイルを閉じてください。')
        return 1




"""ローカル画面。アプリはパスワードを収集・保存しません。"""
import os
import queue
import threading
from datetime import date, datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk




def gui_main():
    root = tk.Tk()
    root.title('MITERAS | 工数レポート')
    root.geometry('780x670')
    root.minsize(740, 650)
    root.configure(bg='#F3F6FA')
    root.option_add('*Font', ('Yu Gothic UI', 10))
    style = ttk.Style(root)
    style.theme_use('clam')
    style.configure('App.TFrame', background='#F3F6FA')
    style.configure('Card.TFrame', background='#FFFFFF')
    style.configure('Card.TLabel', background='#FFFFFF', foreground='#23354B', font=('Yu Gothic UI', 10))
    style.configure('Hint.TLabel', background='#FFFFFF', foreground='#708095', font=('Yu Gothic UI', 9))
    style.configure('Section.TLabel', background='#FFFFFF', foreground='#193B59', font=('Yu Gothic UI', 11, 'bold'))
    style.configure('Status.TLabel', background='#FFFFFF', foreground='#193B59', font=('Yu Gothic UI', 12, 'bold'))
    style.configure('Primary.TButton', background='#2457B7', foreground='white', borderwidth=0,
                    padding=(18, 12), font=('Yu Gothic UI', 11, 'bold'))
    style.map('Primary.TButton', background=[('disabled', '#DCE4EE'), ('active', '#1E4898')],
              foreground=[('disabled', '#8A98AA')])
    style.configure('Secondary.TButton', background='#EAF0F7', foreground='#294560', borderwidth=0, padding=(14, 10))
    style.map('Secondary.TButton', background=[('active', '#DDE7F2')], foreground=[('disabled', '#A0ADBB')])
    style.configure('Compact.TButton', background='#EFF4FA', foreground='#294560', borderwidth=0, padding=(12, 5))
    style.configure('TEntry', padding=(10, 9), fieldbackground='#FAFCFF', bordercolor='#D5DFEB')
    style.configure('Report.Horizontal.TProgressbar', background='#2457B7', troughcolor='#E9EFF6',
                    borderwidth=0, lightcolor='#2457B7', darkcolor='#2457B7')
    outer = ttk.Frame(root, style='App.TFrame', padding=(28, 24))
    outer.pack(fill='both', expand=True)
    tk.Label(outer, text='MITERAS', bg='#F3F6FA', fg='#2457B7', font=('Yu Gothic UI', 10, 'bold')).pack(anchor='w')
    tk.Label(outer, text='工数レポート', bg='#F3F6FA', fg='#193B59', font=('Yu Gothic UI', 24, 'bold')).pack(anchor='w', pady=(2, 4))
    tk.Label(outer, text='期間を選んで、プロジェクト別の工数をExcelにまとめます。',
             bg='#F3F6FA', fg='#6D7F93', font=('Yu Gothic UI', 10)).pack(anchor='w', pady=(0, 20))

    today = datetime.now(JST).date()
    start = tk.StringVar(value=today.replace(day=1).isoformat())
    end = tk.StringVar(value=today.isoformat())
    folder = tk.StringVar(value=str(Path(__file__).resolve().parent))
    card = ttk.Frame(outer, style='Card.TFrame', padding=20)
    card.pack(fill='x')
    top = ttk.Frame(card, style='Card.TFrame')
    top.pack(fill='x', pady=(0, 12))
    ttk.Label(top, text='出力設定', style='Section.TLabel').pack(side='left')
    def this_month():
        day = datetime.now(JST).date()
        start.set(day.replace(day=1).isoformat())
        end.set(day.isoformat())
    month_btn = ttk.Button(top, text='今月', style='Compact.TButton', command=this_month)
    month_btn.pack(side='right')
    date_row = ttk.Frame(card, style='Card.TFrame')
    date_row.pack(fill='x')
    date_row.columnconfigure(0, weight=1)
    date_row.columnconfigure(1, weight=1)
    inputs = []
    for col, (title, var) in enumerate([('開始日', start), ('終了日', end)]):
        group = ttk.Frame(date_row, style='Card.TFrame')
        group.grid(row=0, column=col, sticky='ew', padx=(0, 12) if col == 0 else (12, 0))
        ttk.Label(group, text=title, style='Card.TLabel').pack(anchor='w', pady=(0, 5))
        entry = ttk.Entry(group, textvariable=var)
        entry.pack(fill='x')
        inputs.append(entry)
    ttk.Label(card, text='日付の形式：2026-10-01', style='Hint.TLabel').pack(anchor='w', pady=(5, 14))
    ttk.Label(card, text='保存先', style='Card.TLabel').pack(anchor='w', pady=(0, 5))
    destination_row = ttk.Frame(card, style='Card.TFrame')
    destination_row.pack(fill='x')
    destination = ttk.Entry(destination_row, textvariable=folder)
    destination.pack(side='left', fill='x', expand=True)
    inputs.append(destination)
    def choose():
        result = filedialog.askdirectory(parent=root, title='Excelの保存先', initialdir=folder.get())
        if result:
            folder.set(result)
    browse_btn = ttk.Button(destination_row, text='変更', style='Secondary.TButton', command=choose)
    browse_btn.pack(side='right', padx=(10, 0))

    state_card = ttk.Frame(outer, style='Card.TFrame', padding=20)
    state_card.pack(fill='x', pady=(16, 0))
    headline = tk.StringVar(value='準備完了')
    description = tk.StringVar(value='設定を確認して「MITERASを開く」を押してください。')
    counter = tk.StringVar(value='')
    stage_row = ttk.Frame(state_card, style='Card.TFrame')
    stage_row.pack(fill='x')
    stage_label = ttk.Label(stage_row, textvariable=headline, style='Status.TLabel')
    stage_label.pack(side='left')
    ttk.Label(stage_row, textvariable=counter, style='Hint.TLabel').pack(side='right')
    ttk.Label(state_card, textvariable=description, style='Hint.TLabel', wraplength=655).pack(anchor='w', pady=(7, 12))
    bar = ttk.Progressbar(state_card, style='Report.Horizontal.TProgressbar', mode='determinate', maximum=100)
    bar.pack(fill='x', ipady=3)

    action_row = ttk.Frame(outer, style='App.TFrame')
    action_row.pack(fill='x', pady=(18, 10))
    events = queue.Queue()
    ready, cancel = threading.Event(), threading.Event()
    active = {'running': False, 'phase': 'idle', 'path': None, 'destination': None}

    def lock_inputs(locked):
        for widget in inputs + [month_btn, browse_btn]:
            widget.configure(state='disabled' if locked else 'normal')

    def set_status(title, text, color='#193B59'):
        headline.set(title)
        description.set(text)
        stage_label.configure(foreground=color)

    def wait_login(page):
        events.put(('login', None))
        while not ready.is_set():
            if cancel.is_set():
                raise KeyboardInterrupt
            page.wait_for_timeout(150)
        if cancel.is_set():
            raise KeyboardInterrupt

    def worker(first, last, output):
        try:
            result = run(first, last, output, report=lambda *a, **kw: None,
                         login_wait=wait_login, cancelled=cancel,
                         progress=lambda *value: events.put(('progress', value)),
                         saved=lambda path: events.put(('saved', path)))
            events.put(('done', result))
        except KeyboardInterrupt:
            events.put(('cancelled', None))
        except ImportError:
            events.put(('failed', '必要なライブラリがありません。playwrightとopenpyxlのインストールを確認してください。'))
        except RuntimeError as exc:
            events.put(('failed', str(exc)))
        except Exception:
            events.put(('failed', 'ブラウザーの起動またはExcelの保存に失敗しました。保存先と実行環境を確認してください。'))

    def primary_action():
        if active['phase'] == 'login':
            ready.set()
            active['phase'] = 'export'
            primary.configure(text='出力中…', state='disabled')
            set_status('工数を取得しています', '完了までこの画面でお待ちください。')
            return
        if active['running']:
            return
        try:
            first = date.fromisoformat(start.get().strip())
            last = date.fromisoformat(end.get().strip())
        except ValueError:
            set_status('日付を確認してください', '開始日と終了日をYYYY-MM-DD形式で入力してください。', '#A03D28')
            return
        if last < first or (last - first).days > 365:
            set_status('期間を確認してください', '開始日から終了日までの期間は最大366日です。', '#A03D28')
            return
        dest = Path(folder.get()).expanduser()
        if not dest.is_dir():
            set_status('保存先を確認してください', '既存の保存先フォルダーを選んでください。', '#A03D28')
            return
        output = dest / f'miteras_{first}_{last}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx'
        ready.clear()
        cancel.clear()
        active.update(running=True, phase='opening', path=None, destination=dest)
        lock_inputs(True)
        open_file.configure(state='disabled')
        open_folder.configure(state='disabled')
        primary.configure(state='disabled', text='起動中…')
        cancel_button.configure(state='normal')
        counter.set('')
        bar.configure(value=0)
        set_status('ブラウザーを開いています', 'ChromeでMITERASを開きます。必要に応じてログインしてください。')
        threading.Thread(target=worker, args=(first, last, output), daemon=True).start()

    def stop():
        cancel.set()
        active['phase'] = 'cancelling'
        primary.configure(state='disabled')
        cancel_button.configure(state='disabled')
        set_status('中止しています', '現在の処理が終わるまでお待ちください。')

    def open_result(is_folder=False):
        path = active['destination'] if is_folder else active['path']
        if not path:
            return
        try:
            if sys.platform == 'win32':
                os.startfile(str(path))
            else:
                import subprocess
                subprocess.Popen(['open' if sys.platform == 'darwin' else 'xdg-open', str(path)])
        except Exception:
            set_status('ファイルを開けませんでした', f'保存先から直接開いてください：{path}', '#A03D28')

    primary = ttk.Button(action_row, text='MITERASを開く', style='Primary.TButton', command=primary_action)
    primary.pack(side='left', fill='x', expand=True)
    cancel_button = ttk.Button(action_row, text='中止', state='disabled', style='Secondary.TButton', command=stop)
    cancel_button.pack(side='right', padx=(12, 0))
    result_row = ttk.Frame(outer, style='App.TFrame')
    result_row.pack(fill='x')
    open_file = ttk.Button(result_row, text='Excelを開く', style='Compact.TButton', state='disabled', command=open_result)
    open_file.pack(side='left')
    open_folder = ttk.Button(result_row, text='保存先を開く', style='Compact.TButton', state='disabled', command=lambda: open_result(True))
    open_folder.pack(side='left', padx=(10, 0))

    def poll():
        try:
            while True:
                kind, value = events.get_nowait()
                if kind == 'login':
                    if cancel.is_set():
                        continue
                    active['phase'] = 'login'
                    primary.configure(text='ログイン完了・Excel出力', state='normal')
                    set_status('ログインを確認してください', 'ブラウザーで日別詳細とプロジェクトタブを開き、「ログイン完了・Excel出力」を押してください。')
                elif kind == 'progress':
                    done, total, day, status = value
                    if active['phase'] == 'cancelling':
                        continue
                    if day is None:
                        set_status('Excelを作成しています', '取得したデータの集計と書式を整えています。')
                    else:
                        set_status('工数を取得しています', f'{day:%Y/%m/%d} まで取得しました。')
                    counter.set(f'{done} / {total}日')
                    bar.configure(value=95 * done / max(total, 1))
                elif kind == 'saved':
                    active['path'] = value
                elif kind in ('done', 'failed', 'cancelled'):
                    active.update(running=False, phase='idle')
                    lock_inputs(False)
                    primary.configure(text='もう一度出力する', state='normal')
                    cancel_button.configure(state='disabled')
                    if kind == 'done':
                        bar.configure(value=100)
                        if value == 0:
                            set_status('Excelを出力しました', '「工数明細」の差分と照合結果を確認してください。', '#22744D')
                        else:
                            set_status('一部のデータを取得できませんでした', '未完了のExcelを保存しました。エラーの日を確認してください。', '#A56A16')
                        open_file.configure(state='normal' if active['path'] else 'disabled')
                        open_folder.configure(state='normal')
                    elif kind == 'cancelled':
                        set_status('中止しました', '期間を確認して、再度出力できます。')
                    else:
                        set_status('出力できませんでした', value, '#A03D28')
        except queue.Empty:
            pass
        root.after(120, poll)

    def close():
        if active['running']:
            messagebox.showinfo('処理中', '「中止」を押し、処理が終わってから閉じてください。', parent=root)
            return
        root.destroy()
    root.protocol('WM_DELETE_WINDOW', close)
    poll()
    root.mainloop()


if __name__ == "__main__":
    if "--cli" in sys.argv:
        sys.argv.remove("--cli")
        sys.exit(main())
    gui_main()
